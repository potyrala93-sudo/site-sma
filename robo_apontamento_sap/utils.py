"""
utils.py — Funções de apoio do robô de apontamento (ZFLPP005).

Contém:
- exceções usadas para decidir se uma falha deve ou não ser retentada;
- conversões de dados da planilha (serviço, quantidade, data);
- espera inteligente (esperar_ate) — nada de sleeps fixos;
- tecla de emergência (Ctrl+Shift+Q);
- relatório de execução (CSV gravado linha a linha + Excel no final);
- configuração de logging.
"""
from __future__ import annotations

import csv
import logging
import re
import sys
import threading
import time
import unicodedata
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any, Callable

import yaml

log = logging.getLogger("robo")


# ---------------------------------------------------------------------------
# Exceções
# ---------------------------------------------------------------------------
class ErroNegocio(Exception):
    """Falha esperada do processo (OS não encontrada, serviço divergente...).
    NÃO é retentada: tentar de novo daria o mesmo resultado."""


class ErroSAP(ErroNegocio):
    """O SAP devolveu mensagem de erro (barra de status tipo E/A, popup de erro)."""


class ErroConfiguracao(Exception):
    """Configuração errada (ID inexistente, coluna não encontrada...).
    Interrompe o robô inteiro, pois todas as linhas falhariam igual."""


class TempoEsgotado(Exception):
    """Timeout de espera. Considerado transitório → pode ser retentado."""


class ParadaSolicitada(Exception):
    """Usuário acionou a tecla de emergência."""


# ---------------------------------------------------------------------------
# Configuração e logging
# ---------------------------------------------------------------------------
def carregar_yaml(caminho: str | Path) -> dict:
    with open(caminho, encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def configurar_logging(pasta: str | Path, nivel: str = "INFO") -> Path:
    pasta = Path(pasta)
    pasta.mkdir(parents=True, exist_ok=True)
    arquivo = pasta / f"robo_{datetime.now():%Y%m%d_%H%M%S}.log"
    fmt = logging.Formatter("%(asctime)s [%(levelname)s] %(name)s: %(message)s", "%d/%m/%Y %H:%M:%S")

    raiz = logging.getLogger()
    raiz.setLevel(logging.DEBUG)
    raiz.handlers.clear()

    fh = logging.FileHandler(arquivo, encoding="utf-8")
    fh.setLevel(logging.DEBUG)  # arquivo sempre detalhado
    fh.setFormatter(fmt)
    raiz.addHandler(fh)

    ch = logging.StreamHandler(sys.stdout)
    ch.setLevel(getattr(logging, str(nivel).upper(), logging.INFO))
    ch.setFormatter(fmt)
    raiz.addHandler(ch)
    return arquivo


# ---------------------------------------------------------------------------
# Conversões de dados
# ---------------------------------------------------------------------------
def vazio(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def normalizar_texto(t: Any) -> str:
    """Minúsculas, sem acento, sem espaços duplicados e sem pontuação final.
    Usado para comparar títulos de colunas/popups ('% Exec.' == '% exec')."""
    s = unicodedata.normalize("NFKD", str(t or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"\s+", " ", s).strip().casefold()
    return s.strip(" .:")


def texto_inteiro(v: Any) -> str:
    """Converte valores numéricos vindos da planilha (24001761.0, '24001761,0')
    em texto inteiro, preservando o que já for texto."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip().replace(" ", "").replace(" ", "")
    if re.fullmatch(r"\d+[.,]0+", s):
        s = re.split(r"[.,]", s)[0]
    return s


def normalizar_servico(v: Any) -> str:
    """Normaliza o nº do serviço para comparação, IGNORANDO zeros à esquerda:
    '0024001761' -> '24001761';  24001761.0 -> '24001761'."""
    s = texto_inteiro(v)
    if not s:
        return ""
    return s.lstrip("0") or "0"


def parse_numero(v: Any) -> float:
    """Lê número em formato brasileiro ou americano.
    '45,62' / '45.62' / '1.234,5' / '1,234.5' / 45.62 -> float.
    Regra: se houver '.' e ',', o ÚLTIMO separador é o decimal; se houver só ',',
    ela é decimal; se houver só '.', ele é decimal."""
    if isinstance(v, bool):
        raise ValueError(f"valor não numérico: {v!r}")
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or "").strip().replace(" ", "").replace(" ", "")
    if not s:
        raise ValueError("valor vazio")
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    return float(s)


def formatar_quantidade(v: Any, casas: int = 3) -> str:
    """45.62 -> '45,620' (vírgula decimal, N casas, arredondamento comercial)."""
    n = Decimal(str(parse_numero(v)))
    q = n.quantize(Decimal(1).scaleb(-casas), rounding=ROUND_HALF_UP)
    return f"{q:.{casas}f}".replace(".", ",")


_EPOCA_SHEETS = date(1899, 12, 30)  # dia 0 do número serial do Google Sheets/Excel
_FORMATOS_DATA = (
    "%d/%m/%Y", "%d/%m/%y", "%d.%m.%Y", "%d-%m-%Y", "%Y-%m-%d",
    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d%m%Y",
)


def converter_data(v: Any) -> date:
    """Aceita 'DD/MM/AAAA', datetime/date ou número serial do Google Sheets."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return _EPOCA_SHEETS + timedelta(days=int(v))
    s = str(v or "").strip()
    if not s:
        raise ValueError("data vazia")
    # Serial vindo como texto (ex.: '46296' ou '46296,5'); 8 dígitos = DDMMAAAA
    if re.fullmatch(r"\d{1,6}([.,]\d+)?", s):
        return _EPOCA_SHEETS + timedelta(days=int(float(s.replace(",", "."))))
    for fmt in _FORMATOS_DATA:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"data em formato desconhecido: {s!r}")


def formatar_data(v: Any, fmt: str) -> str:
    return converter_data(v).strftime(fmt)


def primeiro_ultimo_dia(ref: date) -> tuple[date, date]:
    inicio = ref.replace(day=1)
    proximo = (inicio + timedelta(days=32)).replace(day=1)
    return inicio, proximo - timedelta(days=1)


# ---------------------------------------------------------------------------
# Espera inteligente
# ---------------------------------------------------------------------------
def esperar_ate(condicao: Callable[[], Any], timeout: float, intervalo: float = 0.25,
                descricao: str = "condição") -> Any:
    """Chama `condicao()` até ela retornar algo "verdadeiro" ou estourar o timeout.
    Exceções dentro da condição contam como "ainda não" (o SAP rejeita chamadas
    COM enquanto está processando), exceto parada e erro de configuração."""
    fim = time.monotonic() + timeout
    ultimo_erro = None
    while True:
        try:
            r = condicao()
            if r:
                return r
        except (ParadaSolicitada, ErroConfiguracao):
            raise
        except Exception as e:  # noqa: BLE001
            ultimo_erro = e
        if time.monotonic() >= fim:
            extra = f" (último erro: {ultimo_erro})" if ultimo_erro else ""
            raise TempoEsgotado(f"tempo esgotado ({timeout:.0f}s) aguardando {descricao}{extra}")
        time.sleep(intervalo)


# ---------------------------------------------------------------------------
# Tecla de emergência: Ctrl+Shift+Q
# ---------------------------------------------------------------------------
class ParadaEmergencia:
    """Thread que observa o teclado (GetAsyncKeyState) e sinaliza a parada.
    O robô consulta `verificar()` entre as etapas, então ele para num ponto seguro
    (nunca entre o check verde e o Salvar Guia — ver SapClient.checar_parada)."""

    def __init__(self) -> None:
        self._evento = threading.Event()
        self._thread: threading.Thread | None = None

    def iniciar(self) -> None:
        try:
            import win32api  # noqa: F401
        except ImportError:
            log.warning("pywin32 indisponível: tecla de emergência desativada.")
            return
        self._thread = threading.Thread(target=self._observar, daemon=True, name="emergencia")
        self._thread.start()
        log.info("Tecla de emergência ativa: Ctrl+Shift+Q para parar com segurança.")

    def _observar(self) -> None:
        import win32api
        import win32con
        while not self._evento.is_set():
            ctrl = win32api.GetAsyncKeyState(win32con.VK_CONTROL) & 0x8000
            shift = win32api.GetAsyncKeyState(win32con.VK_SHIFT) & 0x8000
            q = win32api.GetAsyncKeyState(ord("Q")) & 0x8000
            if ctrl and shift and q:
                log.warning(">>> Ctrl+Shift+Q detectado: o robô vai parar no próximo ponto seguro.")
                self._evento.set()
            time.sleep(0.1)

    @property
    def acionada(self) -> bool:
        return self._evento.is_set()

    def verificar(self) -> None:
        if self._evento.is_set():
            raise ParadaSolicitada("parada solicitada pelo usuário (Ctrl+Shift+Q)")


# ---------------------------------------------------------------------------
# Relatório de execução
# ---------------------------------------------------------------------------
class Relatorio:
    """Grava cada resultado imediatamente em CSV (não se perde nada se o robô cair)
    e gera um Excel consolidado ao final."""

    CAMPOS = ["linha", "os", "servico", "quantidade", "data_execucao", "status",
              "mensagem", "tentativas", "data_hora"]

    def __init__(self, pasta: str | Path, simulacao: bool) -> None:
        self.pasta = Path(pasta)
        self.pasta.mkdir(parents=True, exist_ok=True)
        sufixo = "_SIMULACAO" if simulacao else ""
        base = f"relatorio_{datetime.now():%Y%m%d_%H%M%S}{sufixo}"
        self.csv = self.pasta / f"{base}.csv"
        self.xlsx = self.pasta / f"{base}.xlsx"
        self.registros: list[dict] = []
        with open(self.csv, "w", newline="", encoding="utf-8-sig") as f:
            csv.DictWriter(f, self.CAMPOS, delimiter=";").writeheader()

    def registrar(self, **campos: Any) -> None:
        campos.setdefault("data_hora", datetime.now().strftime("%d/%m/%Y %H:%M:%S"))
        reg = {c: campos.get(c, "") for c in self.CAMPOS}
        self.registros.append(reg)
        with open(self.csv, "a", newline="", encoding="utf-8-sig") as f:
            csv.DictWriter(f, self.CAMPOS, delimiter=";").writerow(reg)

    def salvar_excel(self) -> Path | None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            log.warning("openpyxl não instalado: relatório só em CSV.")
            return None
        wb = Workbook()
        ws = wb.active
        ws.title = "Execução"
        ws.append(self.CAMPOS)
        for c in ws[1]:
            c.font = Font(bold=True)
        cores = {"OK": "C6EFCE", "ERRO": "FFC7CE", "VERIFICAR": "FFC7CE",
                 "PULADO": "FFEB9C", "SIMULACAO": "DDEBF7", "INTERROMPIDO": "D9D9D9"}
        for r in self.registros:
            ws.append([r[c] for c in self.CAMPOS])
            cor = cores.get(str(r["status"]).split(":")[0])
            if cor:
                ws.cell(ws.max_row, self.CAMPOS.index("status") + 1).fill = PatternFill("solid", fgColor=cor)
        for col in ws.columns:
            largura = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, largura + 2), 80)
        ws.freeze_panes = "A2"
        wb.save(self.xlsx)
        return self.xlsx

    def resumo(self) -> dict[str, int]:
        cont: dict[str, int] = {}
        for r in self.registros:
            chave = str(r["status"]).split(":")[0]
            cont[chave] = cont.get(chave, 0) + 1
        return cont
